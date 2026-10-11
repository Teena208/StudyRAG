from pptx import Presentation

ppt = Presentation("UNIT5_Cybersecurity Principles and Best Practices.pptx")

chunks = []

for slide_number, slide in enumerate(ppt.slides, start=1):
    slide_text = []

    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            slide_text.append(shape.text.strip())

    if slide_text:
        chunk = "\n".join(slide_text)

        chunks.append({
            "slide": slide_number,
            "text": chunk
        })

print(f"Total chunks: {len(chunks)}")

for chunk in chunks:
    print(f"\n--- Chunk from Slide {chunk['slide']} ---")
    print(chunk["text"])