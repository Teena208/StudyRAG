from pptx import Presentation
from sentence_transformers import SentenceTransformer
import chromadb

# Load the PowerPoint
ppt = Presentation("UNIT5_Cybersecurity Principles and Best Practices.pptx")

# Load the embedding model
model = SentenceTransformer("all-MiniLM-L6-v2")

chunks = []

# Extract text from each slide
for slide_number, slide in enumerate(ppt.slides, start=1):
    slide_text = []

    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            slide_text.append(shape.text.strip())

    if slide_text:
        chunks.append({
            "slide": slide_number,
            "text": "\n".join(slide_text)
        })

print("Total chunks:", len(chunks))

# Get the text from each chunk
texts = [chunk["text"] for chunk in chunks]

# Create embeddings
embeddings = model.encode(texts)

print("Total embeddings:", len(embeddings))
print("Size of each embedding:", len(embeddings[0]))

# Create a local ChromaDB database
client = chromadb.PersistentClient(path="./chroma_db")

# Create a collection
collection = client.get_or_create_collection(name="study_material")

# Store chunks, embeddings and slide numbers
collection.add(
    ids=[f"slide_{chunk['slide']}" for chunk in chunks],
    documents=texts,
    embeddings=embeddings.tolist(),
    metadatas=[{"slide": chunk["slide"]} for chunk in chunks]
)

print("Successfully stored chunks in ChromaDB!")
print("Documents in database:", collection.count())